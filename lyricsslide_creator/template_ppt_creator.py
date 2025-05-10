"""
템플릿 PPT 파일을 사용한 가사 슬라이드 생성 모듈

이 모듈은 미리 디자인된 템플릿 PPT 파일을 사용하여 
가사 슬라이드를 생성하는 기능을 제공합니다.
"""

import os
import sys
import shutil
import traceback
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from io import BytesIO


def get_template_path():
    """템플릿 PPT 파일 경로를 반환합니다."""
    # 개발 환경: 프로젝트 루트 경로
    base_path = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    bundled_template = os.path.join(base_path, "template.pptx")
    print(f"템플릿 파일 경로 확인: {bundled_template}")
    
    # 배포 환경: 실행 파일 위치
    if hasattr(sys, 'frozen'):  # PyInstaller로 패키징된 경우
        base_path = os.path.dirname(sys.executable)
        bundled_template = os.path.join(base_path, "template.pptx")
        print(f"PyInstaller 환경에서 템플릿 파일 경로: {bundled_template}")
    
    if os.path.exists(bundled_template):
        print(f"템플릿 파일 발견: {bundled_template}")
        return bundled_template
    
    # 추가 경로 검색 (_internal 폴더 추가)
    possible_paths = [
        bundled_template,
        os.path.join(os.path.dirname(base_path), "template.pptx"),     # 한 단계 상위
        os.path.join(base_path, "input_ppts", "template.pptx"),        # input_ppts 폴더
        os.path.join(os.getcwd(), "template.pptx"),                    # 현재 작업 디렉토리
        os.path.join(base_path, "_internal", "template.pptx"),         # _internal 폴더
        os.path.join(os.path.dirname(base_path), "_internal", "template.pptx")  # 상위/_internal 폴더
    ]
    
    # PyInstaller의 _MEIPASS 경로 확인 (임시 실행 디렉토리)
    if hasattr(sys, '_MEIPASS'):
        possible_paths.append(os.path.join(sys._MEIPASS, "template.pptx"))
        possible_paths.append(os.path.join(sys._MEIPASS, "_internal", "template.pptx"))
        print(f"PyInstaller _MEIPASS 경로 확인: {sys._MEIPASS}")
    
    for path in possible_paths:
        if os.path.exists(path):
            print(f"템플릿 파일을 다른 경로에서 발견: {path}")
            return path
    
    print("템플릿 파일을 어느 위치에서도 찾을 수 없습니다.")
    print(f"현재 확인한 경로들: {possible_paths}")
    
    # 디렉토리 내용 출력 (디버깅용)
    if hasattr(sys, 'frozen'):
        exe_dir = os.path.dirname(sys.executable)
        print(f"실행 파일 디렉토리 내용: {os.listdir(exe_dir)}")
        
        # _internal 폴더가 있으면 내용 출력
        internal_dir = os.path.join(exe_dir, "_internal")
        if os.path.exists(internal_dir) and os.path.isdir(internal_dir):
            print(f"_internal 디렉토리 내용: {os.listdir(internal_dir)}")
    
    return None


def get_font_size_from_user():
    """사용자로부터 폰트 크기를 입력받습니다. 입력이 없으면 기본값 45를 사용합니다."""
    try:
        user_input = input("\n원하는 폰트 크기를 입력하세요 (기본값 45, Enter 키 입력 시 기본값 사용): ")
        if not user_input.strip():
            print("기본 폰트 크기 45pt를 사용합니다.")
            # 항상 새로운 값을 반환하도록 직접 45 반환
            return 45
            
        font_size = int(user_input)
        if font_size <= 0 or font_size > 200:
            print("유효한 범위(1-200)가 아닙니다. 기본 폰트 크기 45pt를 사용합니다.")
            return 45
            
        print(f"폰트 크기 {font_size}pt를 사용합니다.")
        return font_size
    except ValueError:
        print("숫자가 아닌 값이 입력되었습니다. 기본 폰트 크기 45pt를 사용합니다.")
        return 45


def create_lyrics_slides_from_template(output_path, lyrics_pages, song_breaks=None, template_path=None, font_size=45):
    """
    템플릿 PPT 파일을 사용해서 가사 슬라이드를 생성합니다.
    
    Args:
        output_path (str): 출력할 PPT 파일 경로
        lyrics_pages (list): 가사 페이지 목록 (각 페이지는 한 슬라이드에 표시될 텍스트)
        song_breaks (list, optional): 노래 경계 인덱스 목록
        template_path (str, optional): 템플릿 PPT 파일 경로
        font_size (int or list, optional): 폰트 크기 또는 노래별 폰트 크기 목록
    """
    # 템플릿 파일 경로 확인
    if not template_path:
        template_path = get_template_path()
    
    # 템플릿 파일이 없으면 기본 형식으로 생성
    if not template_path or not os.path.exists(template_path):
        print("템플릿 파일을 찾을 수 없습니다. 기본 형식으로 생성합니다.")
        return create_default_lyrics_slides(output_path, lyrics_pages, song_breaks, font_size)
    
    try:
        # 템플릿 파일 복사 (원본 보존)
        temp_path = output_path + ".temp"
        shutil.copy2(template_path, temp_path)
        
        # 복사된 템플릿 열기
        prs = Presentation(temp_path)
        
        # 템플릿의 첫 번째 슬라이드 가져오기
        if len(prs.slides) == 0:
            # 슬라이드가 없으면 빈 레이아웃으로 추가
            template_slide = prs.slides.add_slide(prs.slide_layouts[6])
        else:
            template_slide = prs.slides[0]
        
        # 템플릿 슬라이드에서 텍스트 상자 찾기
        template_text_shapes = []
        for shape in template_slide.shapes:
            if hasattr(shape, 'text_frame'):
                template_text_shapes.append(shape)
        
        # 텍스트 상자가 없으면 기본 형식으로 생성
        if not template_text_shapes:
            print("템플릿에 텍스트 상자가 없습니다. 기본 형식으로 생성합니다.")
            os.remove(temp_path)
            return create_default_lyrics_slides(output_path, lyrics_pages, song_breaks, font_size)
        
        # 기존 슬라이드 모두 삭제
        slide_ids = list(range(len(prs.slides)))
        for idx in reversed(slide_ids):
            rId = prs.slides._sldIdLst[idx].rId
            prs.part.drop_rel(rId)
            del prs.slides._sldIdLst[idx]
        
        # 현재 폰트 크기 초기화 (기본값 또는 첫 곡의 값)
        current_font_size = font_size if isinstance(font_size, int) else font_size[0]
        current_song_index = 0
        
        # 가사 페이지별로 슬라이드 생성
        for i, lyrics in enumerate(lyrics_pages):
            # 노래 구분점 확인 및 폰트 크기 업데이트
            if song_breaks and i in song_breaks:
                # 다음 곡의 인덱스
                current_song_index = song_breaks.index(i) + 1  # +1 because we're on the next song
                
                # 다음 곡에 대한 폰트 크기 설정
                if isinstance(font_size, list) and current_song_index < len(font_size):
                    current_font_size = font_size[current_song_index]
                else:
                    # 설정된 폰트 크기가 없으면 기본값 45 사용
                    current_font_size = 45
                
                print(f"곡 변경: 폰트 크기를 {current_font_size}pt로 설정합니다.")
                
                # 빈 검은색 슬라이드 추가
                blank_slide = prs.slides.add_slide(prs.slide_layouts[6])
                background = blank_slide.background
                fill = background.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(0, 0, 0)
            
            # 새 슬라이드 생성 (비어있는 상태로)
            new_slide = prs.slides.add_slide(prs.slide_layouts[6])
            
            # 배경 설정
            background = new_slide.background
            fill = background.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor(0, 0, 0)  # 검은색 배경
            
            # 비-텍스트 요소 복제 (배경, 이미지 등)
            for shape in template_slide.shapes:
                if not hasattr(shape, 'text_frame'):  # 텍스트 상자가 아닌 요소만 복사
                    try:
                        if hasattr(shape, 'shape_type'):
                            # 도형 복사
                            if shape.shape_type == MSO_SHAPE.AUTO_SHAPE:
                                new_shape = new_slide.shapes.add_shape(
                                    shape.auto_shape_type,
                                    shape.left,
                                    shape.top,
                                    shape.width,
                                    shape.height
                                )
                                
                                # 도형 스타일 복사
                                if hasattr(shape, 'fill'):
                                    new_shape.fill.solid()
                                    if hasattr(shape.fill, 'fore_color') and shape.fill.fore_color:
                                        new_shape.fill.fore_color.rgb = shape.fill.fore_color.rgb
                                
                                # 선 스타일 복사
                                if hasattr(shape, 'line'):
                                    if hasattr(shape.line, 'color') and shape.line.color:
                                        new_shape.line.color.rgb = shape.line.color.rgb
                                    if hasattr(shape.line, 'width'):
                                        new_shape.line.width = shape.line.width
                            
                            # 이미지 복사
                            elif shape.shape_type == MSO_SHAPE.PICTURE:
                                if hasattr(shape, 'image'):
                                    new_slide.shapes.add_picture(
                                        BytesIO(shape.image.blob),
                                        shape.left,
                                        shape.top,
                                        shape.width,
                                        shape.height
                                    )
                    except Exception as e:
                        print(f"요소 복사 중 오류: {e}")
            
            # 가사 텍스트 추가 (첫 번째 텍스트 상자 사용)
            if template_text_shapes:
                shape = template_text_shapes[0]
                
                # 원본 텍스트 상자와 동일한 위치/크기에 새 텍스트 상자 생성
                new_text_box = new_slide.shapes.add_textbox(
                    shape.left,
                    shape.top,
                    shape.width,
                    shape.height
                )
                new_text_frame = new_text_box.text_frame
                
                # 텍스트 프레임 속성 복사
                new_text_frame.word_wrap = True
                if hasattr(shape.text_frame, 'vertical_anchor'):
                    new_text_frame.vertical_anchor = shape.text_frame.vertical_anchor
                if hasattr(shape.text_frame, 'margin_left'):
                    new_text_frame.margin_left = shape.text_frame.margin_left
                if hasattr(shape.text_frame, 'margin_right'):
                    new_text_frame.margin_right = shape.text_frame.margin_right
                if hasattr(shape.text_frame, 'margin_top'):
                    new_text_frame.margin_top = shape.text_frame.margin_top
                if hasattr(shape.text_frame, 'margin_bottom'):
                    new_text_frame.margin_bottom = shape.text_frame.margin_bottom
                
                # 가사 텍스트 추가
                p = new_text_frame.add_paragraph()
                p.text = lyrics
                
                # 폰트 스타일 설정 - 현재 곡의 폰트 크기 사용
                p.font.size = Pt(current_font_size)
                p.font.name = "맑은 고딕"
                p.font.bold = True
                p.alignment = PP_ALIGN.CENTER
                
                # 원본 텍스트 색상 유지 시도
                try:
                    if shape.text_frame.paragraphs and hasattr(shape.text_frame.paragraphs[0], 'font'):
                        orig_font = shape.text_frame.paragraphs[0].font
                        if hasattr(orig_font, 'color') and orig_font.color and orig_font.color.rgb:
                            p.font.color.rgb = orig_font.color.rgb
                        else:
                            p.font.color.rgb = RGBColor(255, 255, 255)  # 기본 흰색
                    else:
                        p.font.color.rgb = RGBColor(255, 255, 255)  # 기본 흰색
                except:
                    p.font.color.rgb = RGBColor(255, 255, 255)  # 기본 흰색
        
        # PPT 저장
        prs.save(output_path)
        
        # 임시 파일 삭제
        try:
            os.remove(temp_path)
        except Exception:
            pass
            
        print(f"템플릿 기반 PPT 생성 완료: {output_path}")
        return output_path
        
    except Exception as e:
        print(f"템플릿 기반 PPT 생성 중 오류 발생: {str(e)}")
        traceback.print_exc()
        
        # 임시 파일 정리
        try:
            if 'temp_path' in locals() and os.path.exists(temp_path):
                os.remove(temp_path)
        except:
            pass
            
        # 오류 발생 시 기본 형식으로 생성
        return create_default_lyrics_slides(output_path, lyrics_pages, song_breaks, font_size)


def create_default_lyrics_slides(output_path, lyrics_pages, song_breaks=None, font_size=45):
    """
    템플릿 없이 기본 형식으로 가사 슬라이드를 생성합니다.
    
    Args:
        output_path (str): 출력할 PPT 파일 경로
        lyrics_pages (list): 가사 페이지 목록
        song_breaks (list, optional): 노래 경계 인덱스 목록
        font_size (int or list, optional): 폰트 크기 또는 노래별 폰트 크기 목록
    """
    try:
        # 새 프레젠테이션 생성
        prs = Presentation()
        
        # 슬라이드 크기 설정 (16:9 비율)
        prs.slide_width = Inches(16)
        prs.slide_height = Inches(9)
        
        # 현재 폰트 크기 초기화
        current_font_size = font_size if isinstance(font_size, int) else font_size[0]
        current_song_index = 0
        
        # 각 가사 페이지별 슬라이드 생성
        for i, lyrics in enumerate(lyrics_pages):
            # 노래 구분점 확인 및 폰트 크기 업데이트
            if song_breaks and i in song_breaks:
                # 다음 곡의 인덱스
                current_song_index = song_breaks.index(i) + 1
                
                # 다음 곡에 대한 폰트 크기 설정
                if isinstance(font_size, list) and current_song_index < len(font_size):
                    current_font_size = font_size[current_song_index]
                else:
                    # 설정된 폰트 크기가 없으면 기본값 45 사용
                    current_font_size = 45
                
                print(f"곡 변경: 폰트 크기를 {current_font_size}pt로 설정합니다.")
                
                # 빈 검은색 슬라이드 추가
                blank_slide = prs.slides.add_slide(prs.slide_layouts[6])
                background = blank_slide.background
                fill = background.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(0, 0, 0)
            
            # 빈 슬라이드 추가
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            
            # 배경을 검은색으로 설정
            background = slide.background
            fill = background.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor(0, 0, 0)
            
            # 슬라이드 크기 가져오기
            slide_width = prs.slide_width
            slide_height = prs.slide_height
            
            # 텍스트 상자 추가 - 가사용 텍스트 상자만 추가
            left = 0
            top = Inches(1.5)
            width = slide_width
            height = Inches(4.0)
            
            text_box = slide.shapes.add_textbox(left, top, width, height)
            text_frame = text_box.text_frame
            text_frame.word_wrap = True
            text_frame.vertical_anchor = MSO_ANCHOR.TOP
            
            # 가사 텍스트 추가
            p = text_frame.add_paragraph()
            p.text = lyrics
            p.font.name = "맑은 고딕"
            p.font.size = Pt(current_font_size)  # 현재 폰트 크기 사용
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.CENTER
        
        # PPT 저장
        prs.save(output_path)
        print(f"기본 형식 PPT 생성 완료: {output_path}")
        return output_path
        
    except Exception as e:
        print(f"기본 PPT 생성 중 오류 발생: {str(e)}")
        traceback.print_exc()
        return None