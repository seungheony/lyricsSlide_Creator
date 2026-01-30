from pptx import Presentation
import copy
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from lxml import etree
import io

class PPTBuilder:
    def __init__(self, template_path):
        self.template_path = template_path
        self.prs = None
        try:
            self.prs = Presentation(template_path)
        except Exception as e:
            print(f"Failed to load template: {e}")
            raise

    def duplicate_slide_elements(self, source_slide, target_slide):
        for shape in source_slide.shapes:
            el = shape.element
            new_el = copy.deepcopy(el)
            target_slide.shapes._spTree.insert_element_before(new_el, 'p:extLst')

    def replace_text_in_slide(self, slide, lyrics_chunk):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            if '가사' not in shape.text:
                continue

            text_frame = shape.text_frame
            
            try:
                original_p = text_frame.paragraphs[0]
                original_run = original_p.runs[0]
                original_font = copy.deepcopy(original_run.font)
                original_alignment = original_p.alignment
            except IndexError:
                continue

            # 첫 번째 문단을 제외한 나머지 문단을 삭제
            for p in text_frame.paragraphs[1:]:
                p._p.getparent().remove(p._p)

            # 첫 번째 문단의 내용을 모두 지움
            p = text_frame.paragraphs[0]
            p.clear()
            p.alignment = original_alignment

            lines = lyrics_chunk.split('\n')
            for i, line in enumerate(lines):
                if i > 0:
                    p = text_frame.add_paragraph()
                    p.alignment = original_alignment

                run = p.add_run()
                run.text = line
                
                run.font.name = original_font.name
                run.font.size = original_font.size
                run.font.bold = original_font.bold
                run.font.italic = original_font.italic
                if original_font.color and hasattr(original_font.color, 'rgb') and original_font.color.rgb is not None:
                    run.font.color.rgb = original_font.color.rgb
            
            break

    def generate_slides_to_memory(self, lyrics_data):
        if not self.prs or not self.prs.slides: return None

        template_slide = self.prs.slides[0]
        slide_layout = template_slide.slide_layout

        for i in range(len(self.prs.slides) - 1, 0, -1):
            rId = self.prs.slides._sldIdLst[i].rId
            self.prs.part.drop_rel(rId)
            del self.prs.slides._sldIdLst[i]

        for chunk in lyrics_data:
            new_slide = self.prs.slides.add_slide(slide_layout)
            for shape in new_slide.shapes:
                 if shape.is_placeholder:
                    sp = shape.element
                    sp.getparent().remove(sp)
            self.duplicate_slide_elements(template_slide, new_slide)
            self.replace_text_in_slide(new_slide, chunk)

        self.delete_template_slide(0)
        
        ppt_buffer = io.BytesIO()
        self.prs.save(ppt_buffer)
        ppt_buffer.seek(0)
        return ppt_buffer

    def delete_template_slide(self, index):
        try:
            rId = self.prs.slides._sldIdLst[index].rId
            self.prs.part.drop_rel(rId)
            del self.prs.slides._sldIdLst[index]
        except: pass

    @staticmethod
    def save_to_file(buffer, output_path):
        try:
            with open(output_path, 'wb') as f:
                f.write(buffer.read())
            return True
        except Exception:
            return False

    @staticmethod
    def split_lyrics(full_lyrics, lines_per_slide=2):
        lines = [l.strip() for l in full_lyrics.split('\n') if l.strip()]
        return ["\n".join(lines[i:i+lines_per_slide]) for i in range(0, len(lines), lines_per_slide)]